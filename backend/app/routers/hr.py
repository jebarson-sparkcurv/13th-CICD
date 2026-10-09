"""Staff Leave & Attendance (module key ``leave_attendance``).

Staff self-service (SiteEngineer / Accountant / ProcurementOfficer):
  GET  /me/leave                       balances + my requests
  POST /me/leave                       apply for leave
  POST /me/leave/{id}/cancel           cancel my pending request
  GET  /me/attendance?month=YYYY-MM    my month + today's record
  POST /me/attendance/check-in
  POST /me/attendance/check-out

Admin:
  GET  /leave/policies                 PUT /leave/policies
  GET  /leave/requests                 list (filters: status, user_id, month)
  POST /leave/requests/{id}/approve    POST /leave/requests/{id}/reject
  POST /leave/requests/{id}/cancel     (undo an approved leave)
  GET  /staff-attendance?month=        monthly grid
  PUT  /staff-attendance               set one day for one person
  DELETE /staff-attendance?user_id=&date=
  POST /staff-attendance/mark-all      mark everyone unmarked as present for a date
  GET  /staff-attendance/export?month= Excel

Rules
  * Sundays are the weekly off: never counted as leave days.
  * Approved leave writes "leave" attendance rows (locked from manual edits).
  * Payroll deducts: absent = 1 day, half_day = 0.5, unpaid leave = 1 (0.5 if half).
"""
import io
from calendar import monthrange
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal
from typing import Optional
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import User, Notification
from app.models.hr import (LeavePolicy, LeaveRequest, StaffAttendance, STAFF_ROLES,
                           DEFAULT_POLICIES, ATTENDANCE_STATUSES)
from app.core.security import get_current_user, _effective_modules

router = APIRouter(tags=["leave & attendance"])

TZ = ZoneInfo("Asia/Kolkata")
MODULE = "leave_attendance"
MAX_LEAVE_SPAN = 60          # days in one request
MAX_BACKDATE = 30            # days in the past a request may start


# --------------------------------------------------------------------------
# Access
# --------------------------------------------------------------------------

def _module_on(user: User) -> bool:
    return user.role == "SuperAdmin" or MODULE in (_effective_modules(user) or [])


def staff_user(user: User = Depends(get_current_user)) -> User:
    if user.role not in STAFF_ROLES:
        raise HTTPException(status_code=403, detail="Leave & attendance is for staff (Site Engineer, Accountant, Procurement Officer)")
    if not _module_on(user):
        raise HTTPException(status_code=403, detail="Leave & Attendance is not enabled for your account")
    return user


def hr_admin(user: User = Depends(get_current_user)) -> User:
    if user.role == "SuperAdmin":
        raise HTTPException(status_code=403, detail="Use “Login as Admin” for the company to manage its leave & attendance")
    if user.role != "Admin":
        raise HTTPException(status_code=403, detail="Only Admins can do this")
    if not _module_on(user):
        raise HTTPException(status_code=403, detail="Leave & Attendance is not enabled for your company")
    return user


def _tid(user: User) -> int:
    return user.tenant_id or 1


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------

def today_local() -> date:
    return datetime.now(TZ).date()


def working_dates(start: date, end: date) -> list[date]:
    out, d = [], start
    while d <= end:
        if d.weekday() != 6:          # Sunday off
            out.append(d)
        d += timedelta(days=1)
    return out


def _parse_date(v, field: str) -> date:
    if isinstance(v, date):
        return v
    try:
        return date.fromisoformat(str(v)[:10])
    except (TypeError, ValueError):
        raise HTTPException(status_code=422, detail=f"Invalid {field}")


def _parse_month(month: Optional[str]) -> tuple[date, date]:
    if not month:
        t = today_local()
        y, m = t.year, t.month
    else:
        try:
            y, m = (int(x) for x in month.split("-")[:2])
            date(y, m, 1)
        except (TypeError, ValueError):
            raise HTTPException(status_code=422, detail="month must be YYYY-MM")
    return date(y, m, 1), date(y, m, monthrange(y, m)[1])


def _hhmm(d: date, v) -> Optional[datetime]:
    if v in (None, ""):
        return None
    try:
        hh, mm = (int(x) for x in str(v).split(":")[:2])
        return datetime.combine(d, time(hh, mm), tzinfo=TZ)
    except (TypeError, ValueError):
        raise HTTPException(status_code=422, detail="Time must be HH:MM")


def _local_hhmm(dt: Optional[datetime]) -> Optional[str]:
    return dt.astimezone(TZ).strftime("%H:%M") if dt else None


def ensure_policies(db: Session, tenant_id: int) -> list[LeavePolicy]:
    pols = db.query(LeavePolicy).filter(LeavePolicy.tenant_id == tenant_id).all()
    if not pols:
        for t, q, paid in DEFAULT_POLICIES:
            db.add(LeavePolicy(tenant_id=tenant_id, leave_type=t,
                               annual_quota=Decimal(q) if q is not None else None, is_paid=paid))
        db.commit()
        pols = db.query(LeavePolicy).filter(LeavePolicy.tenant_id == tenant_id).all()
    order = {t: i for i, (t, _, _) in enumerate(DEFAULT_POLICIES)}
    return sorted(pols, key=lambda p: (order.get(p.leave_type, 99), p.leave_type))


def _policy(db: Session, tenant_id: int, leave_type: str) -> LeavePolicy:
    p = next((x for x in ensure_policies(db, tenant_id) if x.leave_type == leave_type and x.is_active), None)
    if not p:
        raise HTTPException(status_code=422, detail=f"Unknown leave type: {leave_type}")
    return p


def _used(db: Session, user_id: int, leave_type: str, year: int, statuses, exclude_id=None) -> float:
    q = db.query(func.coalesce(func.sum(LeaveRequest.days), 0)).filter(
        LeaveRequest.user_id == user_id, LeaveRequest.leave_type == leave_type,
        LeaveRequest.status.in_(statuses),
        LeaveRequest.start_date >= date(year, 1, 1), LeaveRequest.start_date <= date(year, 12, 31))
    if exclude_id:
        q = q.filter(LeaveRequest.id != exclude_id)
    return float(q.scalar() or 0)


def balances(db: Session, user: User, year: int) -> list[dict]:
    out = []
    for p in ensure_policies(db, _tid(user)):
        if not p.is_active:
            continue
        used = _used(db, user.id, p.leave_type, year, ["Approved"])
        pending = _used(db, user.id, p.leave_type, year, ["Pending"])
        quota = float(p.annual_quota) if p.annual_quota is not None else None
        out.append({"leave_type": p.leave_type, "is_paid": p.is_paid, "quota": quota,
                    "used": used, "pending": pending,
                    "remaining": None if quota is None else round(quota - used - pending, 1)})
    return out


def _notify(db: Session, user_ids, ntype: str, title: str, message: str, link: str):
    for uid in {u for u in user_ids if u}:
        db.add(Notification(user_id=uid, type=ntype, title=title, message=message, link=link))


def _admins(db: Session, tenant_id: int) -> list[int]:
    return [u.id for u in db.query(User).filter(User.tenant_id == tenant_id, User.role == "Admin",
                                                 User.status != "Disabled").all()]


def _staff_q(db: Session, user: User):
    return (db.query(User).filter(User.role.in_(STAFF_ROLES), User.status != "Disabled",
                                  User.tenant_id == user.tenant_id).order_by(User.name))


def _get_request(db: Session, rid: int, user: User) -> LeaveRequest:
    r = db.get(LeaveRequest, rid)
    if not r or r.tenant_id != user.tenant_id:
        raise HTTPException(status_code=404, detail="Leave request not found")
    return r


def request_out(r: LeaveRequest) -> dict:
    return {"id": r.id, "user_id": r.user_id, "user_name": r.user.name if r.user else None,
            "user_role": r.user.role if r.user else None,
            "leave_type": r.leave_type, "start_date": r.start_date.isoformat(),
            "end_date": r.end_date.isoformat(), "half_day": r.half_day, "days": float(r.days),
            "reason": r.reason, "status": r.status,
            "decided_by_name": r.decider.name if r.decider else None,
            "decided_at": r.decided_at.isoformat() if r.decided_at else None,
            "decision_note": r.decision_note,
            "created_at": r.created_at.isoformat() if r.created_at else None}


def att_out(a: StaffAttendance) -> dict:
    return {"id": a.id, "user_id": a.user_id, "date": a.attendance_date.isoformat(),
            "status": a.status, "check_in": _local_hhmm(a.check_in), "check_out": _local_hhmm(a.check_out),
            "source": a.source, "leave_type": a.leave_type, "is_half": a.is_half, "note": a.note}


def _unpaid_types(db: Session, tenant_id: int) -> set[str]:
    return {p.leave_type for p in ensure_policies(db, tenant_id) if not p.is_paid}


def deduction_days(rows, unpaid_types: set[str]) -> tuple[float, dict]:
    counts = {"absent": 0.0, "half_day": 0.0, "unpaid_leave": 0.0}
    for a in rows:
        if a.attendance_date.weekday() == 6:     # weekly off is never deducted
            continue
        if a.status == "absent":
            counts["absent"] += 1
        elif a.status == "half_day":
            counts["half_day"] += 1
        elif a.status == "leave" and a.leave_type in unpaid_types:
            counts["unpaid_leave"] += 0.5 if a.is_half else 1
    total = counts["absent"] + 0.5 * counts["half_day"] + counts["unpaid_leave"]
    return total, counts


def breakdown_from_rows(rows_by_date: dict, unpaid: set[str], start: date, end: date,
                        as_of: date, attendance_based: bool = True) -> dict:
    """Pure calculation — attendance-based pay for one person.

    Counts days from ``start`` up to ``as_of`` (inclusive):
      paid       present, approved paid leave, Sundays, week-off, holiday
      deducted   absent (1), unpaid leave (1, ½ if half-day), half_day status (½),
                 and UNMARKED working days (treated as absent)
    Days after ``as_of`` are ignored (month not finished yet).
    """
    period_days = (end - start).days + 1
    out = {"period_days": period_days, "working_days": 0, "present": 0.0, "half_day": 0,
           "paid_leave": 0.0, "unpaid_leave": 0.0, "absent": 0, "unmarked": 0,
           "deduct_days": 0.0, "attendance_based": attendance_based,
           "counted_until": as_of.isoformat() if as_of >= start else None}
    if not attendance_based:
        return out
    d = start
    while d <= min(end, as_of):
        a = rows_by_date.get(d)
        if d.weekday() == 6:
            pass                                    # Sunday = weekly off: always paid
        elif a is None:
            out["working_days"] += 1
            out["unmarked"] += 1                    # nobody marked it → not paid
        else:
            out["working_days"] += 1
            half = 0.5 if a.is_half else 1.0
            if a.status == "absent":
                out["absent"] += 1
            elif a.status == "half_day":
                out["half_day"] += 1
                out["present"] += 0.5
            elif a.status == "leave":
                if a.leave_type in unpaid:
                    out["unpaid_leave"] += half
                else:
                    out["paid_leave"] += half
                if a.is_half:
                    out["present"] += 0.5           # worked the other half
            elif a.status == "present":
                out["present"] += 1
        d += timedelta(days=1)
    out["deduct_days"] = out["absent"] + out["unmarked"] + out["unpaid_leave"] + 0.5 * out["half_day"]
    return out


def pay_breakdown(db: Session, user: User, start: date, end: date, final: bool = False) -> dict:
    """``final=True`` (processing a finished month) counts every day through
    ``end``; otherwise only days before today are counted."""
    today = today_local()
    as_of = end if final else min(end, today - timedelta(days=1))
    if user.role not in STAFF_ROLES:
        return breakdown_from_rows({}, set(), start, end, as_of, attendance_based=False)
    rows = {a.attendance_date: a for a in db.query(StaffAttendance).filter(
        StaffAttendance.user_id == user.id, StaffAttendance.attendance_date >= start,
        StaffAttendance.attendance_date <= end).all()}
    return breakdown_from_rows(rows, _unpaid_types(db, user.tenant_id or 1), start, end, as_of)


def pay_note(b: dict) -> str:
    parts = []
    for k, label in (("absent", "absent"), ("unmarked", "unmarked"), ("half_day", "half day"),
                     ("unpaid_leave", "unpaid leave")):
        if b.get(k):
            parts.append(f"{b[k]:g} {label}")
    return ", ".join(parts)


def payroll_deduction(db: Session, user: User, start: date, end: date) -> tuple[float, str]:
    b = pay_breakdown(db, user, start, end)
    return b["deduct_days"], pay_note(b)


# --------------------------------------------------------------------------
# Staff self-service — leave
# --------------------------------------------------------------------------

@router.get("/me/leave")
def my_leave(db: Session = Depends(get_db), user: User = Depends(staff_user)):
    year = today_local().year
    reqs = (db.query(LeaveRequest).filter(LeaveRequest.user_id == user.id)
            .order_by(LeaveRequest.start_date.desc(), LeaveRequest.id.desc()).limit(100).all())
    return {"year": year, "balances": balances(db, user, year), "requests": [request_out(r) for r in reqs]}


@router.post("/me/leave", status_code=201)
def apply_leave(body: dict, db: Session = Depends(get_db), user: User = Depends(staff_user)):
    tid = _tid(user)
    pol = _policy(db, tid, body.get("leave_type") or "")
    start = _parse_date(body.get("start_date"), "start date")
    end = _parse_date(body.get("end_date") or body.get("start_date"), "end date")
    half = bool(body.get("half_day"))
    reason = (str(body.get("reason") or "").strip())[:2000] or None
    today = today_local()
    if end < start:
        raise HTTPException(status_code=422, detail="End date is before start date")
    if (end - start).days + 1 > MAX_LEAVE_SPAN:
        raise HTTPException(status_code=422, detail=f"One request can cover at most {MAX_LEAVE_SPAN} days")
    if start < today - timedelta(days=MAX_BACKDATE):
        raise HTTPException(status_code=422, detail=f"Leave can't start more than {MAX_BACKDATE} days ago")
    if start.year != end.year:
        raise HTTPException(status_code=422, detail="Leave across the new year: please send one request per year")
    if start > today + timedelta(days=365):
        raise HTTPException(status_code=422, detail="Leave can be requested up to one year ahead")
    if half and start != end:
        raise HTTPException(status_code=422, detail="Half-day leave must be a single date")
    work = working_dates(start, end)
    if not work:
        raise HTTPException(status_code=422, detail="Those dates are all Sundays (weekly off)")
    days = 0.5 if half else float(len(work))
    overlap = db.query(LeaveRequest).filter(
        LeaveRequest.user_id == user.id, LeaveRequest.status.in_(["Pending", "Approved"]),
        LeaveRequest.start_date <= end, LeaveRequest.end_date >= start).first()
    if overlap:
        raise HTTPException(status_code=409, detail=f"You already have a {overlap.status.lower()} request "
                                                    f"for {overlap.start_date} – {overlap.end_date}")
    if pol.annual_quota is not None:
        taken = _used(db, user.id, pol.leave_type, start.year, ["Approved", "Pending"])
        left = float(pol.annual_quota) - taken
        if days > left:
            raise HTTPException(status_code=422, detail=f"Not enough {pol.leave_type} leave: {max(left, 0):g} day(s) left, "
                                                        f"this request needs {days:g}. Try Unpaid leave.")
    r = LeaveRequest(tenant_id=tid, user_id=user.id, leave_type=pol.leave_type, start_date=start,
                     end_date=end, half_day=half, days=Decimal(str(days)), reason=reason, status="Pending")
    db.add(r); db.flush()
    span = f"{start:%d %b}" + ("" if start == end else f" – {end:%d %b}")
    _notify(db, _admins(db, tid), "LeaveRequested", f"Leave request: {user.name}",
            f"{pol.leave_type} · {span} · {days:g} day(s)", "/admin/users?tab=leave")
    db.commit(); db.refresh(r)
    return request_out(r)


@router.post("/me/leave/{rid}/cancel")
def cancel_my_leave(rid: int, db: Session = Depends(get_db), user: User = Depends(staff_user)):
    r = db.get(LeaveRequest, rid)
    if not r or r.user_id != user.id:
        raise HTTPException(status_code=404, detail="Leave request not found")
    if r.status != "Pending":
        raise HTTPException(status_code=409, detail="Only pending requests can be cancelled — ask your Admin")
    r.status = "Cancelled"
    db.commit(); db.refresh(r)
    return request_out(r)


# --------------------------------------------------------------------------
# Staff self-service — attendance
# --------------------------------------------------------------------------

def _row(db: Session, user_id: int, d: date) -> Optional[StaffAttendance]:
    return db.query(StaffAttendance).filter(StaffAttendance.user_id == user_id,
                                            StaffAttendance.attendance_date == d).first()


@router.get("/me/attendance")
def my_attendance(month: Optional[str] = None, db: Session = Depends(get_db),
                  user: User = Depends(staff_user)):
    start, end = _parse_month(month)
    rows = (db.query(StaffAttendance).filter(StaffAttendance.user_id == user.id,
                                             StaffAttendance.attendance_date >= start,
                                             StaffAttendance.attendance_date <= end)
            .order_by(StaffAttendance.attendance_date).all())
    today = _row(db, user.id, today_local())
    summary = {s: 0 for s in ATTENDANCE_STATUSES}
    for a in rows:
        summary[a.status] = summary.get(a.status, 0) + 1
    return {"month": f"{start:%Y-%m}", "today": att_out(today) if today else None,
            "today_date": today_local().isoformat(), "records": [att_out(a) for a in rows],
            "summary": summary}


@router.post("/me/attendance/check-in")
def check_in(db: Session = Depends(get_db), user: User = Depends(staff_user)):
    d = today_local()
    now = datetime.now(TZ)
    a = _row(db, user.id, d)
    if a and a.status == "leave" and not a.is_half:
        raise HTTPException(status_code=409, detail="You're on approved leave today")
    if a and a.check_in:
        raise HTTPException(status_code=409, detail=f"Already checked in at {_local_hhmm(a.check_in)}")
    if not a:
        a = StaffAttendance(tenant_id=_tid(user), user_id=user.id, attendance_date=d,
                            status="present", source="self", marked_by=user.id)
        db.add(a)
    elif a.status in ("absent", "week_off", "holiday"):
        a.status, a.source = "present", "self"
    a.check_in = now
    try:
        db.commit()
    except IntegrityError:          # double-click: the other request already checked in
        db.rollback()
        raise HTTPException(status_code=409, detail="Already checked in")
    db.refresh(a)
    return att_out(a)


@router.post("/me/attendance/check-out")
def check_out(db: Session = Depends(get_db), user: User = Depends(staff_user)):
    a = _row(db, user.id, today_local())
    if not a or not a.check_in:
        raise HTTPException(status_code=409, detail="Check in first")
    if a.check_out:
        raise HTTPException(status_code=409, detail=f"Already checked out at {_local_hhmm(a.check_out)}")
    a.check_out = datetime.now(TZ)
    db.commit(); db.refresh(a)
    return att_out(a)


# --------------------------------------------------------------------------
# Admin — leave policies & requests
# --------------------------------------------------------------------------

@router.get("/leave/policies")
def get_policies(db: Session = Depends(get_db), user: User = Depends(hr_admin)):
    return [{"leave_type": p.leave_type, "annual_quota": float(p.annual_quota) if p.annual_quota is not None else None,
             "is_paid": p.is_paid, "is_active": p.is_active} for p in ensure_policies(db, _tid(user))]


@router.put("/leave/policies")
def put_policies(body: list[dict], db: Session = Depends(get_db), user: User = Depends(hr_admin)):
    tid = _tid(user)
    existing = {p.leave_type: p for p in ensure_policies(db, tid)}
    seen = set()
    for item in body:
        name = str(item.get("leave_type") or "").strip()[:40]
        if not name:
            raise HTTPException(status_code=422, detail="Leave type name is required")
        if name.lower() in seen:
            raise HTTPException(status_code=422, detail=f"Leave type “{name}” is listed twice")
        seen.add(name.lower())
        q = item.get("annual_quota")
        if q in ("", None):
            quota = None
        else:
            try:
                quota = Decimal(str(q))
            except Exception:  # noqa: BLE001
                raise HTTPException(status_code=422, detail=f"Quota for {name} must be a number")
            if not quota.is_finite() or quota < 0 or quota > 365:
                raise HTTPException(status_code=422, detail=f"Quota for {name} must be 0–365")
        p = existing.get(name) or LeavePolicy(tenant_id=tid, leave_type=name)
        p.annual_quota = quota
        p.is_paid = bool(item.get("is_paid", True))
        p.is_active = bool(item.get("is_active", True))
        if p.id is None:
            db.add(p)
    db.commit()
    return get_policies(db=db, user=user)


@router.get("/leave/requests")
def list_requests(status: Optional[str] = None, user_id: Optional[int] = None,
                  month: Optional[str] = None, db: Session = Depends(get_db),
                  user: User = Depends(hr_admin)):
    q = db.query(LeaveRequest).filter(LeaveRequest.tenant_id == user.tenant_id)
    if status:
        q = q.filter(LeaveRequest.status.in_(status.split(",")))
    if user_id:
        q = q.filter(LeaveRequest.user_id == user_id)
    if month:
        s, e = _parse_month(month)
        q = q.filter(LeaveRequest.start_date <= e, LeaveRequest.end_date >= s)
    rows = q.order_by((LeaveRequest.status == "Pending").desc(), LeaveRequest.start_date.desc(),
                      LeaveRequest.id.desc()).limit(500).all()
    out = []
    bal_cache: dict = {}
    for r in rows:
        o = request_out(r)
        key = (r.user_id, r.start_date.year)
        if key not in bal_cache and r.user:
            bal_cache[key] = {b["leave_type"]: b for b in balances(db, r.user, r.start_date.year)}
        b = bal_cache.get(key, {}).get(r.leave_type)
        o["balance_remaining"] = b["remaining"] if b else None
        out.append(o)
    pending = sum(1 for r in rows if r.status == "Pending")
    return {"pending_count": pending, "requests": out}


def _write_leave_rows(db: Session, r: LeaveRequest, admin: User):
    for d in working_dates(r.start_date, r.end_date):
        a = _row(db, r.user_id, d)
        if not a:
            a = StaffAttendance(tenant_id=r.tenant_id, user_id=r.user_id, attendance_date=d)
            db.add(a)
        elif a.source != "leave":
            a.prev_status, a.prev_source = a.status, a.source
        a.status, a.source = "leave", "leave"
        a.leave_request_id, a.leave_type, a.is_half = r.id, r.leave_type, r.half_day
        a.marked_by = admin.id


def _clear_leave_rows(db: Session, r: LeaveRequest):
    for a in db.query(StaffAttendance).filter(StaffAttendance.leave_request_id == r.id).all():
        if a.prev_status:   # day was marked before the leave — put it back
            a.status, a.source = a.prev_status, a.prev_source or "manual"
        elif a.check_in:    # they worked that day — keep it as present
            a.status, a.source = "present", "self"
        else:
            db.delete(a)
            continue
        a.leave_request_id, a.leave_type, a.is_half = None, None, False
        a.prev_status = a.prev_source = None


@router.post("/leave/requests/{rid}/approve")
def approve(rid: int, body: Optional[dict] = None, db: Session = Depends(get_db),
            user: User = Depends(hr_admin)):
    r = _get_request(db, rid, user)
    if r.status != "Pending":
        raise HTTPException(status_code=409, detail=f"Request is already {r.status}")
    pol = _policy(db, r.tenant_id, r.leave_type)
    if pol.annual_quota is not None:
        approved = _used(db, r.user_id, r.leave_type, r.start_date.year, ["Approved"])
        if approved + float(r.days) > float(pol.annual_quota):
            raise HTTPException(status_code=422, detail=f"Approving would exceed the {r.leave_type} quota "
                                                        f"({float(pol.annual_quota) - approved:g} day(s) left)")
    r.status, r.decided_by, r.decided_at = "Approved", user.id, datetime.now(timezone.utc)
    r.decision_note = (str((body or {}).get("note") or "").strip())[:1000] or None
    _write_leave_rows(db, r, user)
    _notify(db, [r.user_id], "LeaveApproved", "Leave approved",
            f"{r.leave_type} · {r.start_date:%d %b} – {r.end_date:%d %b}", "/my/attendance")
    db.commit(); db.refresh(r)
    return request_out(r)


@router.post("/leave/requests/{rid}/reject")
def reject(rid: int, body: Optional[dict] = None, db: Session = Depends(get_db),
           user: User = Depends(hr_admin)):
    r = _get_request(db, rid, user)
    if r.status != "Pending":
        raise HTTPException(status_code=409, detail=f"Request is already {r.status}")
    r.status, r.decided_by, r.decided_at = "Rejected", user.id, datetime.now(timezone.utc)
    r.decision_note = (str((body or {}).get("note") or "").strip())[:1000] or None
    _notify(db, [r.user_id], "LeaveRejected", "Leave not approved",
            (r.decision_note or f"{r.leave_type} · {r.start_date:%d %b} – {r.end_date:%d %b}"), "/my/attendance")
    db.commit(); db.refresh(r)
    return request_out(r)


@router.post("/leave/requests/{rid}/cancel")
def admin_cancel(rid: int, db: Session = Depends(get_db), user: User = Depends(hr_admin)):
    r = _get_request(db, rid, user)
    if r.status not in ("Pending", "Approved"):
        raise HTTPException(status_code=409, detail=f"Request is already {r.status}")
    if r.status == "Approved":
        _clear_leave_rows(db, r)
    r.status, r.decided_by, r.decided_at = "Cancelled", user.id, datetime.now(timezone.utc)
    _notify(db, [r.user_id], "LeaveCancelled", "Leave cancelled by Admin",
            f"{r.leave_type} · {r.start_date:%d %b} – {r.end_date:%d %b}", "/my/attendance")
    db.commit(); db.refresh(r)
    return request_out(r)


# --------------------------------------------------------------------------
# Admin — attendance grid
# --------------------------------------------------------------------------

def _grid(db: Session, user: User, month: Optional[str]):
    start, end = _parse_month(month)
    staff = _staff_q(db, user).all()
    ids = [u.id for u in staff] or [0]
    rows = db.query(StaffAttendance).filter(StaffAttendance.user_id.in_(ids),
                                            StaffAttendance.attendance_date >= start,
                                            StaffAttendance.attendance_date <= end).all()
    unpaid = _unpaid_types(db, _tid(user))
    by_user: dict = {u.id: {} for u in staff}
    for a in rows:
        by_user.setdefault(a.user_id, {})[a.attendance_date.isoformat()] = a
    today = today_local()
    days = []
    d = start
    while d <= end:
        days.append({"date": d.isoformat(), "day": d.day, "weekday": d.strftime("%a"),
                     "is_sunday": d.weekday() == 6, "is_future": d > today})
        d += timedelta(days=1)
    people = []
    for u in staff:
        recs = by_user.get(u.id, {})
        c = {s: 0 for s in ATTENDANCE_STATUSES}
        for a in recs.values():
            c[a.status] = c.get(a.status, 0) + 1
        b = breakdown_from_rows({a.attendance_date: a for a in recs.values()}, unpaid, start, end,
                                min(end, today - timedelta(days=1)))   # same rule payroll uses
        ded = b["deduct_days"]
        unmarked = b["unmarked"]
        people.append({"id": u.id, "name": u.name, "role": u.role,
                       "base_salary": float(u.base_salary) if getattr(u, "base_salary", None) else None,
                       "records": {k: att_out(a) for k, a in recs.items()},
                       "summary": {**c, "unmarked": unmarked, "deduction_days": ded}})
    return start, end, days, people


@router.get("/staff-attendance")
def attendance_grid(month: Optional[str] = None, db: Session = Depends(get_db),
                    user: User = Depends(hr_admin)):
    start, _, days, people = _grid(db, user, month)
    return {"month": f"{start:%Y-%m}", "today": today_local().isoformat(), "days": days, "staff": people}


def _staff_member(db: Session, uid, user: User) -> User:
    try:
        u = db.get(User, int(uid))
    except (TypeError, ValueError):
        u = None
    if not u or u.role not in STAFF_ROLES or u.tenant_id != user.tenant_id:
        raise HTTPException(status_code=404, detail="Staff member not found")
    return u


@router.put("/staff-attendance")
def set_attendance(body: dict, db: Session = Depends(get_db), user: User = Depends(hr_admin)):
    u = _staff_member(db, body.get("user_id"), user)
    d = _parse_date(body.get("date"), "date")
    if d > today_local():
        raise HTTPException(status_code=422, detail="Can't mark attendance for a future date")
    status = body.get("status")
    if status == "leave":
        raise HTTPException(status_code=422, detail="Leave is set by approving a leave request")
    if status not in ATTENDANCE_STATUSES:
        raise HTTPException(status_code=422, detail=f"Invalid status: {status}")
    if d.weekday() == 6 and status in ("absent", "half_day"):
        raise HTTPException(status_code=422, detail="Sunday is the weekly off — it can't be marked absent")
    a = _row(db, u.id, d)
    if a and a.source == "leave":
        raise HTTPException(status_code=409, detail="This day is approved leave — cancel the leave request to change it")
    if not a:
        a = StaffAttendance(tenant_id=u.tenant_id or 1, user_id=u.id, attendance_date=d)
        db.add(a)
    a.status, a.source, a.marked_by = status, "manual", user.id
    if "check_in" in body:
        a.check_in = _hhmm(d, body.get("check_in"))
    if "check_out" in body:
        a.check_out = _hhmm(d, body.get("check_out"))
    if a.check_in and a.check_out and a.check_out < a.check_in:
        raise HTTPException(status_code=422, detail="Check-out is before check-in")
    if "note" in body:
        a.note = (str(body.get("note") or "").strip())[:500] or None
    db.commit(); db.refresh(a)
    return att_out(a)


@router.delete("/staff-attendance")
def clear_attendance(user_id: int, date_: str = Query(..., alias="date"),
                     db: Session = Depends(get_db), user: User = Depends(hr_admin)):
    u = _staff_member(db, user_id, user)
    a = _row(db, u.id, _parse_date(date_, "date"))
    if a and a.source == "leave":
        raise HTTPException(status_code=409, detail="This day is approved leave — cancel the leave request to change it")
    if a:
        db.delete(a); db.commit()
    return {"status": "cleared"}


@router.post("/staff-attendance/mark-all")
def mark_all_present(body: dict, db: Session = Depends(get_db), user: User = Depends(hr_admin)):
    d = _parse_date(body.get("date"), "date")
    if d > today_local():
        raise HTTPException(status_code=422, detail="Can't mark attendance for a future date")
    n = 0
    for u in _staff_q(db, user).all():
        if _row(db, u.id, d):
            continue
        db.add(StaffAttendance(tenant_id=u.tenant_id or 1, user_id=u.id, attendance_date=d,
                               status="present", source="manual", marked_by=user.id))
        n += 1
    db.commit()
    return {"marked": n}


@router.get("/staff-attendance/export")
def export_attendance(month: Optional[str] = None, db: Session = Depends(get_db),
                      user: User = Depends(hr_admin)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    start, _, days, people = _grid(db, user, month)
    code = {"present": "P", "absent": "A", "half_day": "½", "leave": "L", "week_off": "WO", "holiday": "H"}
    fill = {"P": "D1FAE5", "A": "FEE2E2", "½": "FEF3C7", "L": "DBEAFE", "WO": "F1F5F9", "H": "F1F5F9"}
    wb = Workbook()
    ws = wb.active
    ws.title = f"{start:%b %Y}"
    ws["A1"] = f"Staff attendance — {start:%B %Y}"
    ws["A1"].font = Font(bold=True, size=14)
    header = ["Staff", "Role"] + [f"{x['day']}\n{x['weekday'][:2]}" for x in days] + \
             ["Present", "Absent", "Half day", "Leave", "Unmarked", "Deduct days"]
    for i, h in enumerate(header, start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for r, p in enumerate(people, start=4):
        ws.cell(row=r, column=1, value=p["name"])
        ws.cell(row=r, column=2, value=p["role"])
        for j, x in enumerate(days, start=3):
            rec = p["records"].get(x["date"])
            v = code.get(rec["status"]) if rec else ("WO" if x["is_sunday"] else "")
            if rec and rec["status"] == "leave" and rec["is_half"]:
                v = "L½"
            c = ws.cell(row=r, column=j, value=v)
            c.alignment = Alignment(horizontal="center")
            if v:
                c.fill = PatternFill("solid", fgColor=fill.get(v.rstrip("½") or v, "FFFFFF"))
        s = p["summary"]
        for k, v in enumerate([s["present"], s["absent"], s["half_day"], s["leave"], s["unmarked"],
                               s["deduction_days"]], start=3 + len(days)):
            ws.cell(row=r, column=k, value=v)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 18
    for j in range(3, 3 + len(days)):
        ws.column_dimensions[ws.cell(row=3, column=j).column_letter].width = 4.5
    ws.row_dimensions[3].height = 30
    note_row = 5 + len(people)
    ws.cell(row=note_row, column=1, value="P present · A absent · ½ half day · L leave · L½ half-day leave · WO weekly off · H holiday")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="staff-attendance-{start:%Y-%m}.xlsx"'})
